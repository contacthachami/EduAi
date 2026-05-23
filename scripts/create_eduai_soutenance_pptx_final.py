from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "EduAI_Soutenance_19_20_FINAL.pptx"
FIG = ROOT / "rapport" / "figures"
DOCS = ROOT / "docs"


SLIDE_W = 13.333
SLIDE_H = 7.5


def C(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip("#")
    return RGBColor(
        int(hex_value[0:2], 16),
        int(hex_value[2:4], 16),
        int(hex_value[4:6], 16),
    )


# Palette derived from the real frontend theme + one controlled AI blue.
PAPER = C("#F6F3EF")
SUBTLE = C("#FAF8F5")
CARD = C("#FFFFFF")
INK = C("#171412")
INK_2 = C("#5F5750")
MUTED = C("#8D837A")
BORDER = C("#DED8D1")
BORDER_STRONG = C("#BDB4AA")
COPPER = C("#B85B2A")
COPPER_DARK = C("#93451F")
COPPER_SOFT = C("#F8EFE9")
BLUE = C("#2563EB")
BLUE_SOFT = C("#EFF6FF")
NAVY = C("#1E293B")
NAVY_2 = C("#334155")
GREEN = C("#2F6A4E")
GREEN_SOFT = C("#E5F2EA")
RED = C("#9B2F2F")
RED_SOFT = C("#F8E6E4")
WARN = C("#8A5A16")
WARN_SOFT = C("#F5EAD5")
WHITE = C("#FFFFFF")


def I(x: float):
    return Inches(x)


def fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def line(shape, color=BORDER, width=0.75):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)


def no_line(shape):
    shape.line.fill.background()


def rect(slide, x, y, w, h, color=CARD, border=BORDER, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    fill(s, color)
    if border is None:
        no_line(s)
    else:
        line(s, border)
    return s


def oval(slide, x, y, w, h, color=CARD, border=BORDER):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(x), I(y), I(w), I(h))
    fill(s, color)
    if border is None:
        no_line(s)
    else:
        line(s, border)
    return s


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=14,
    color=INK,
    bold=False,
    font="DM Sans",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_notes(slide, value: str):
    slide.notes_slide.notes_text_frame.text = value.strip()


def add_bg(slide, color=PAPER):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def brand(slide, x=0.55, y=0.35, dark=False):
    rect(slide, x, y, 0.42, 0.42, COPPER, None)
    text(slide, "E", x + 0.11, y + 0.08, 0.2, 0.2, 16, WHITE, True, "DM Sans", PP_ALIGN.CENTER)
    text(slide, "EduAI", x + 0.52, y + 0.08, 1.0, 0.25, 13.5, WHITE if dark else INK, True, "DM Sans")


def footer(slide, n, dark=False):
    col = C("#D7CCC4") if dark else MUTED
    text(slide, "EduAI - Soutenance PFE 2025/2026", 0.6, 7.14, 4.8, 0.18, 8.2, col)
    text(slide, f"{n:02d}/15", 12.35, 7.14, 0.55, 0.18, 8.2, col, align=PP_ALIGN.RIGHT)


def header(slide, section, title, subtitle, n, bg=PAPER):
    add_bg(slide, bg)
    brand(slide)
    text(slide, section.upper(), 0.6, 0.92, 2.7, 0.18, 8.7, COPPER, True, "JetBrains Mono")
    text(slide, title, 0.6, 1.18, 8.8, 0.5, 24, INK, True, "Playfair Display")
    if subtitle:
        text(slide, subtitle, 0.62, 1.72, 8.6, 0.26, 11.5, INK_2)
    footer(slide, n)


def chip(slide, value, x, y, w, color=COPPER, bg=None, size=8.5):
    bg = bg or color
    s = rect(slide, x, y, w, 0.34, bg, None)
    text(slide, value, x + 0.08, y + 0.105, w - 0.16, 0.1, size, WHITE if bg in [COPPER, BLUE, NAVY, GREEN, RED] else color, True, align=PP_ALIGN.CENTER)
    return s


def mini_card(slide, title, body, x, y, w, h, accent=COPPER, value=None, fill_color=CARD):
    rect(slide, x, y, w, h, fill_color, BORDER)
    rect(slide, x, y, 0.07, h, accent, None, radius=False)
    if value:
        text(slide, value, x + 0.25, y + 0.2, w - 0.4, 0.34, 24, accent, True, "DM Sans")
        text(slide, title, x + 0.27, y + 0.76, w - 0.4, 0.2, 10.5, INK, True)
        text(slide, body, x + 0.27, y + 1.06, w - 0.45, h - 1.1, 9.4, INK_2)
    else:
        text(slide, title, x + 0.25, y + 0.18, w - 0.4, 0.24, 12.3, INK, True)
        text(slide, body, x + 0.25, y + 0.54, w - 0.45, h - 0.6, 9.4, INK_2)


def picture_fit(slide, path: Path, x, y, w, h, border=True, bg=CARD):
    rect(slide, x, y, w, h, bg, BORDER if border else None)
    if not path.exists():
        text(slide, path.name, x + 0.15, y + h / 2 - 0.1, w - 0.3, 0.2, 8, MUTED, align=PP_ALIGN.CENTER)
        return
    im = Image.open(path)
    iw, ih = im.size
    aspect = iw / ih
    box_aspect = w / h
    if aspect > box_aspect:
        pic_w = w - 0.22
        pic_h = pic_w / aspect
    else:
        pic_h = h - 0.22
        pic_w = pic_h * aspect
    px = x + (w - pic_w) / 2
    py = y + (h - pic_h) / 2
    slide.shapes.add_picture(str(path), I(px), I(py), width=I(pic_w), height=I(pic_h))


def browser(slide, x, y, w, h, title, accent=COPPER, lines=None, tag=None):
    rect(slide, x, y, w, h, CARD, BORDER)
    rect(slide, x, y, w, 0.36, SUBTLE, BORDER, radius=True)
    for i, c in enumerate([RED, WARN, GREEN]):
        oval(slide, x + 0.15 + i * 0.16, y + 0.14, 0.07, 0.07, c, None)
    text(slide, title, x + 0.52, y + 0.11, w - 0.7, 0.12, 6.7, MUTED)
    rect(slide, x + 0.25, y + 0.62, w - 0.5, 0.42, COPPER_SOFT if accent == COPPER else BLUE_SOFT, None)
    text(slide, title.split(" - ")[0], x + 0.42, y + 0.74, w - 0.84, 0.16, 9.4, INK, True)
    if tag:
        chip(slide, tag, x + w - 1.08, y + 0.68, 0.78, accent, size=6.6)
    lines = lines or ["Module", "Interaction", "Résultat"]
    yy = y + 1.25
    for i, line_value in enumerate(lines[:4]):
        rect(slide, x + 0.32, yy, w - 0.64, 0.26, SUBTLE, BORDER)
        text(slide, line_value, x + 0.45, yy + 0.075, w - 0.9, 0.08, 6.8, INK_2)
        yy += 0.38


def metric(slide, value, label, body, x, y, w, color):
    mini_card(slide, label, body, x, y, w, 1.7, color, value=value)


def slide_1(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, CARD)
    rect(s, 0, 0, 13.333, 7.5, CARD, None, radius=False)
    rect(s, 8.35, 0, 4.98, 7.5, PAPER, None, radius=False)
    rect(s, 8.9, 0.65, 3.75, 5.95, CARD, BORDER)
    browser(s, 9.25, 1.0, 3.05, 1.85, "Workspace - Cours", BLUE, ["Questions RAG", "Résumé", "Quiz"], "READY")
    browser(s, 9.25, 3.15, 3.05, 1.55, "Mode examen", GREEN, ["Timer 24:36", "Score 82%", "Note B"], "EXAM")
    rect(s, 9.25, 5.0, 1.38, 0.74, BLUE_SOFT, None)
    text(s, "RAG", 9.54, 5.24, 0.76, 0.18, 13, BLUE, True, align=PP_ALIGN.CENTER)
    rect(s, 10.9, 5.0, 1.38, 0.74, COPPER_SOFT, None)
    text(s, "SM-2", 11.18, 5.24, 0.76, 0.18, 13, COPPER, True, align=PP_ALIGN.CENTER)

    brand(s, 0.62, 0.55)
    text(s, "Assistant Pédagogique Intelligent", 0.65, 1.45, 5.8, 0.32, 17, COPPER, True)
    text(s, "EduAI", 0.6, 1.88, 5.6, 0.8, 48, INK, True, "Playfair Display")
    text(
        s,
        "Transformer un PDF de cours en espace de compréhension, de révision et d'auto-évaluation.",
        0.65,
        2.78,
        6.25,
        0.62,
        17,
        INK_2,
    )
    rect(s, 0.65, 3.65, 5.85, 0.02, COPPER, None, radius=False)
    for i, line_value in enumerate(
        [
            "Projet de Fin d'Études - 2025/2026",
            "EL Mehdi HACHAMI",
            "Bachelor IA & Sciences des Données",
            "EST Essaouira - Université Cadi Ayyad",
            "Encadrant : Pr. El Mahdi ERRAJI",
            "Entreprise : 2Pi eLearning SARL",
        ]
    ):
        text(s, line_value, 0.68, 4.02 + i * 0.31, 5.9, 0.18, 11.2, INK if i == 1 else INK_2, i == 1)
    picture_fit(s, FIG / "logo_este.png", 0.68, 6.45, 0.85, 0.42, border=False, bg=CARD)
    footer(s, 1)
    add_notes(
        s,
        """
Bonjour Mesdames et Messieurs les membres du jury. Je suis EL Mehdi HACHAMI, étudiant en Bachelor Ingénierie Informatique en Intelligence Artificielle et Sciences des Données. Aujourd'hui, je vais vous présenter mon projet de fin d'études : EduAI, un assistant pédagogique intelligent. L'idée principale est simple : au lieu de laisser l'étudiant seul devant un PDF long et passif, EduAI transforme ce document en un espace interactif pour comprendre, réviser et s'auto-évaluer.
""",
    )


def slide_2(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Constat", "Le PDF ne suffit pas pour apprendre", "Le support est disponible, mais l'apprentissage reste souvent passif.", 2)
    rect(s, 0.8, 2.25, 3.25, 3.75, CARD, BORDER)
    text(s, "PDF de cours", 1.15, 2.58, 2.45, 0.28, 16, INK, True, align=PP_ALIGN.CENTER)
    for i in range(12):
        rect(s, 1.15, 3.05 + i * 0.2, 2.35 if i % 4 else 1.65, 0.055, C("#D8CFC7"), None, radius=False)
    chip(s, "80 pages", 1.43, 5.45, 0.95, RED, size=7.5)
    chip(s, "statique", 2.55, 5.45, 0.95, WARN, size=7.5)
    text(s, "->", 4.48, 3.78, 0.36, 0.22, 18, COPPER, True, align=PP_ALIGN.CENTER)
    problems = [
        ("Compréhension", "L'étudiant n'a pas d'explication adaptée à ses questions.", BLUE),
        ("Révision", "Le résumé, les fiches et les quiz restent à créer manuellement.", COPPER),
        ("Auto-évaluation", "Aucun score clair pour savoir si l'on est prêt.", RED),
    ]
    for i, (title, body, color) in enumerate(problems):
        mini_card(s, title, body, 5.05, 2.15 + i * 1.15, 3.15, 0.9, color)
    rect(s, 8.85, 2.15, 3.85, 3.35, NAVY, None)
    text(s, "Problème central", 9.2, 2.58, 3.15, 0.22, 13, C("#FBD6BF"), True, align=PP_ALIGN.CENTER)
    text(
        s,
        "L'étudiant doit comprendre, mémoriser et mesurer son niveau sans outil interactif adapté.",
        9.25,
        3.15,
        3.05,
        1.15,
        19,
        WHITE,
        True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
    )
    chip(s, "Opportunité : apprentissage actif", 9.35, 4.75, 2.8, COPPER, size=9.2)
    add_notes(
        s,
        """
Le point de départ du projet vient d'une situation très courante dans l'enseignement supérieur marocain. L'étudiant reçoit un PDF de cours, parfois très long, parfois très dense, et il doit se débrouiller seul pour identifier l'essentiel, comprendre les notions, créer ses fiches, préparer des questions et vérifier s'il est prêt pour l'examen. Le PDF est donc utile comme support de distribution, mais il ne suffit pas comme outil d'apprentissage.
""",
    )


def slide_3(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Marché", "Pourquoi maintenant au Maroc ?", "Un besoin massif, un public connecté et un contexte favorable au numérique éducatif.", 3, CARD)
    metric(s, "1,2M", "TAM", "Étudiants dans l'enseignement supérieur au Maroc.", 0.85, 2.08, 3.55, BLUE)
    metric(s, "400k", "SAM", "Étudiants connectés disposant d'un ordinateur ou smartphone.", 4.88, 2.08, 3.55, COPPER)
    metric(s, "5k", "SOM A1", "Objectif réaliste d'utilisateurs visés en première année.", 8.9, 2.08, 3.55, GREEN)
    rect(s, 0.85, 4.55, 11.6, 1.35, SUBTLE, BORDER)
    text(s, "Contexte favorable", 1.2, 4.9, 2.4, 0.22, 14, INK, True)
    text(
        s,
        "Maroc Digital 2030 + usage croissant du numérique éducatif + besoin d'outils accessibles pour les cours réels.",
        3.35,
        4.9,
        8.3,
        0.36,
        14.5,
        INK_2,
    )
    chip(s, "Prix adapté", 3.35, 5.45, 1.25, COPPER, size=8.2)
    chip(s, "PDF réels", 4.85, 5.45, 1.15, BLUE, size=8.2)
    chip(s, "SaaS", 6.25, 5.45, 0.8, GREEN, size=8.2)
    add_notes(
        s,
        """
Ce besoin n'est pas isolé. Le Maroc compte environ 1,2 million d'étudiants dans l'enseignement supérieur. Parmi eux, une population importante est déjà connectée et habituée aux outils numériques. En parallèle, la stratégie Maroc Digital 2030 encourage la transformation numérique, y compris dans l'éducation. Cela crée une opportunité claire : proposer une solution pensée pour les étudiants marocains, avec un prix accessible et une expérience adaptée à leurs vrais supports de cours.
""",
    )


def slide_4(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Vision produit", "De document statique à parcours d'apprentissage", "EduAI couvre tout le cycle : comprendre, réviser, s'auto-évaluer, progresser.", 4)
    stages = [("PDF", "Support de cours"), ("IA", "Traitement intelligent"), ("Workspace", "Outils d'apprentissage")]
    xs = [0.95, 5.05, 9.15]
    colors = [NAVY, BLUE, COPPER]
    for i, (name, desc) in enumerate(stages):
        rect(s, xs[i], 2.25, 3.05, 1.15, CARD, BORDER)
        text(s, name, xs[i] + 0.2, 2.56, 2.65, 0.25, 20, colors[i], True, "Playfair Display", PP_ALIGN.CENTER)
        text(s, desc, xs[i] + 0.25, 3.02, 2.55, 0.15, 9.2, INK_2, align=PP_ALIGN.CENTER)
        if i < 2:
            text(s, "->", xs[i] + 3.35, 2.7, 0.35, 0.2, 17, COPPER, True, align=PP_ALIGN.CENTER)
    actions = [("Comprendre", BLUE), ("Réviser", COPPER), ("S'auto-évaluer", GREEN), ("Progresser", NAVY)]
    for i, (a, c) in enumerate(actions):
        chip(s, a, 1.15 + i * 2.75, 4.2, 2.2, c, size=11)
    modules = [("Résumé", BLUE_SOFT, BLUE), ("Q&A", BLUE_SOFT, BLUE), ("Quiz", COPPER_SOFT, COPPER), ("Flashcards", COPPER_SOFT, COPPER), ("Mindmap", GREEN_SOFT, GREEN), ("Examen", SUBTLE, NAVY)]
    for i, (m, bg, col) in enumerate(modules):
        rect(s, 1.0 + i * 1.95, 5.35, 1.65, 0.52, bg, None)
        text(s, m, 1.12 + i * 1.95, 5.52, 1.4, 0.1, 8.6, col, True, align=PP_ALIGN.CENTER)
    add_notes(
        s,
        """
La vision d'EduAI est de faire évoluer le PDF d'un simple document vers un espace de travail intelligent. L'étudiant importe son cours, et la plateforme génère automatiquement des outils d'apprentissage : résumé, questions-réponses, quiz, flashcards, carte mentale et examen simulé. L'objectif n'est pas seulement de lire plus vite, mais d'apprendre mieux, avec des activités qui poussent l'étudiant à comprendre, mémoriser et mesurer son niveau.
""",
    )


def slide_5(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Produit", "L'expérience utilisateur déjà pensée dans le code", "Landing, auth, upload, workspace et admin : le parcours est complet.", 5, CARD)
    browser(s, 0.75, 2.05, 2.45, 1.7, "Landing", COPPER, ["Créer mon compte", "Pipeline NLP", "Fonctions clés"], "WEB")
    browser(s, 3.45, 2.05, 2.45, 1.7, "Auth", BLUE, ["Login", "Register", "JWT 7 jours"], "JWT")
    browser(s, 6.15, 2.05, 2.45, 1.7, "Upload PDF", COPPER, ["Drag & drop", "PyMuPDF", "Indexation"], "PDF")
    browser(s, 8.85, 2.05, 3.3, 1.7, "Espace cours", GREEN, ["Questions", "Synthèse", "Quiz", "Export PDF"], "READY")
    browser(s, 4.1, 4.28, 5.15, 1.62, "Admin", NAVY, ["Utilisateurs", "Plans Free / Pro", "Stats", "Protection rôle"], "ROLE")
    text(s, "Parcours simple", 0.9, 4.55, 2.55, 0.28, 16, INK, True, "Playfair Display")
    text(s, "Créer un compte -> importer un PDF -> apprendre avec l'IA.", 0.9, 5.0, 2.75, 0.44, 12, INK_2)
    text(s, "Côté jury : montrer vite que ce n'est pas un concept, c'est un produit fonctionnel.", 9.55, 4.78, 2.65, 0.52, 11, COPPER, True)
    add_notes(
        s,
        """
J'ai voulu garder l'expérience très simple. L'étudiant n'a pas besoin de comprendre le fonctionnement interne de l'IA. Il crée un compte, importe son PDF, puis il accède à un espace de cours où tous les outils sont disponibles. La partie administration permet aussi de gérer les utilisateurs, les rôles et les plans Free ou Pro. Cette simplicité est importante, parce qu'une solution éducative doit réduire la charge de travail, pas ajouter une complexité supplémentaire.
""",
    )


def slide_6(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Démonstration", "Comprendre : Q&A RAG + résumé", "L'étudiant pose une question, obtient une réponse ancrée dans son cours et peut vérifier les sources.", 6)
    rect(s, 0.85, 2.05, 5.55, 3.95, CARD, BORDER)
    text(s, "Chat Q&A RAG", 1.15, 2.35, 2.2, 0.24, 15, INK, True)
    chip(s, "Streaming SSE", 4.7, 2.32, 1.2, BLUE, size=7.8)
    rect(s, 1.2, 2.9, 4.35, 0.55, BLUE_SOFT, C("#BFDBFE"))
    text(s, "Quelle est l'idée principale du chapitre ?", 1.42, 3.08, 3.7, 0.12, 9.2, INK)
    rect(s, 1.8, 3.68, 4.05, 1.22, SUBTLE, BORDER)
    text(s, "Réponse pédagogique", 2.05, 3.92, 2.2, 0.18, 11.5, BLUE, True)
    text(s, "Définition -> explication -> exemple -> à retenir", 2.05, 4.28, 3.35, 0.18, 9, INK_2)
    chip(s, "Source p.4", 2.05, 4.63, 0.9, COPPER, size=6.8)
    chip(s, "Source p.7", 3.1, 4.63, 0.9, COPPER, size=6.8)
    rect(s, 6.8, 2.05, 5.65, 3.95, CARD, BORDER)
    text(s, "Résumé par chapitre", 7.1, 2.35, 2.4, 0.24, 15, INK, True)
    chip(s, "Export PDF", 10.85, 2.32, 1.05, COPPER, size=7.8)
    for i, title in enumerate(["Chapitre 1", "Chapitre 2", "Chapitre 3"]):
        y = 2.95 + i * 0.82
        rect(s, 7.2, y, 4.75, 0.56, SUBTLE, BORDER)
        text(s, title, 7.42, y + 0.16, 1.15, 0.12, 8.7, COPPER, True)
        rect(s, 8.55, y + 0.18, 2.9, 0.07, C("#D8CFC7"), None, radius=False)
        rect(s, 8.55, y + 0.36, 2.35, 0.06, C("#E6DED6"), None, radius=False)
    add_notes(
        s,
        """
La première fonctionnalité clé est le chat de questions-réponses. Contrairement à un chatbot généraliste, EduAI ne répond pas uniquement à partir de sa connaissance générale. Il recherche d'abord les passages les plus pertinents dans le cours importé, puis génère une réponse pédagogique à partir de ces passages. Les sources sont affichées pour que l'étudiant puisse vérifier. Le résumé complète cette expérience en identifiant l'essentiel par chapitre, avec une possibilité d'export PDF.
""",
    )


def slide_7(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Démonstration", "Réviser et s'évaluer : apprentissage actif", "Quiz, flashcards, mindmap et examen transforment la révision en activité mesurable.", 7, CARD)
    items = [
        ("Quiz", "Tester", "QCM générés automatiquement\nScore + explication", BLUE),
        ("Flashcards SM-2", "Mémoriser", "Répétition espacée\ninterval, ease_factor", COPPER),
        ("Mindmap", "Relier", "Concepts et relations\nReactFlow", GREEN),
        ("Mode examen", "Mesurer", "Timer, pourcentage,\nnote A-F", NAVY),
    ]
    for i, (title, verb, body, color) in enumerate(items):
        x = 0.75 + i * 3.05
        rect(s, x, 2.05, 2.65, 3.35, CARD, BORDER)
        text(s, title, x + 0.25, 2.35, 2.15, 0.22, 13.5, INK, True)
        chip(s, verb, x + 0.25, 2.78, 1.05, color, size=8.2)
        if title == "Mindmap":
            oval(s, x + 1.02, 3.45, 0.58, 0.58, GREEN_SOFT, C("#BBF7D0"))
            for dx, dy in [(-0.55, -0.28), (0.55, -0.28), (-0.55, 0.55), (0.55, 0.55)]:
                oval(s, x + 1.31 + dx, 3.72 + dy, 0.34, 0.34, CARD, C("#BBF7D0"))
        elif title == "Mode examen":
            text(s, "24:36", x + 0.77, 3.45, 1.0, 0.3, 19, RED, True, align=PP_ALIGN.CENTER)
            text(s, "82%  B", x + 0.75, 4.0, 1.1, 0.22, 16, GREEN, True, align=PP_ALIGN.CENTER)
        else:
            for j in range(4):
                rect(s, x + 0.35, 3.35 + j * 0.28, 1.95, 0.15, COPPER_SOFT if j == 1 else SUBTLE, BORDER)
        text(s, body, x + 0.3, 4.62, 2.05, 0.42, 9.1, INK_2, align=PP_ALIGN.CENTER)
    rect(s, 1.15, 5.9, 11.0, 0.42, NAVY, None)
    text(s, "Objectif : passer d'une lecture passive à une progression visible.", 1.5, 6.05, 10.3, 0.1, 11.5, WHITE, True, align=PP_ALIGN.CENTER)
    add_notes(
        s,
        """
EduAI ne se limite pas à expliquer. Il pousse l'étudiant à réviser activement. Les quiz permettent de tester la compréhension. Les flashcards utilisent l'algorithme SM-2, connu notamment dans Anki, pour organiser la répétition espacée. La carte mentale aide à visualiser les relations entre les concepts. Enfin, le mode examen répond à une question très concrète : est-ce que je suis prêt ? L'étudiant obtient un score, un pourcentage, une note de A à F et des corrections.
""",
    )


def slide_8(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "IA", "Pipeline IA : la vraie valeur technique", "PDF -> extraction -> chunking -> embeddings -> FAISS -> RAG -> génération pédagogique.", 8)
    steps = [
        ("PDF", "cours", NAVY),
        ("PyMuPDF", "extraction", COPPER),
        ("Chunks", "400 + 50", WARN),
        ("MiniLM", "384-d", GREEN),
        ("FAISS", "IndexFlatIP", BLUE),
        ("Top-3", "sources", COPPER),
        ("Llama 3.3", "70B RAG", NAVY),
    ]
    for i, (name, sub, color) in enumerate(steps):
        x = 0.45 + i * 1.84
        rect(s, x, 2.55, 1.47, 0.95, CARD, BORDER)
        text(s, name, x + 0.08, 2.84, 1.3, 0.16, 10.6, color, True, align=PP_ALIGN.CENTER)
        text(s, sub, x + 0.08, 3.16, 1.3, 0.1, 7.2, MUTED, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            text(s, "->", x + 1.5, 2.92, 0.28, 0.15, 13, COPPER, True, align=PP_ALIGN.CENTER)
    facts = [
        ("Top-k", "3 chunks", BLUE),
        ("Seuil", "0.25 / 0.30", COPPER),
        ("Langues", "FR / AR / EN", GREEN),
        ("Fallback", "CamemBERT", NAVY),
    ]
    for i, (k, v, c) in enumerate(facts):
        mini_card(s, k, v, 1.0 + i * 3.0, 4.5, 2.25, 0.82, c)
    rect(s, 1.0, 5.78, 11.35, 0.48, CARD, BORDER)
    text(s, "Pourquoi c'est fort : les réponses sont contextualisées, vérifiables et réutilisées pour résumé, quiz, flashcards, mindmap et examen.", 1.25, 5.96, 10.8, 0.12, 9.5, INK_2, True, align=PP_ALIGN.CENTER)
    add_notes(
        s,
        """
Le coeur technique d'EduAI est ce pipeline. Lors de l'upload, le PDF est lu avec PyMuPDF, puis découpé en chunks de 400 tokens avec un overlap de 50 pour préserver le contexte. Chaque chunk est transformé en vecteur de 384 dimensions grâce au modèle multilingue MiniLM. Ces vecteurs sont indexés dans FAISS avec IndexFlatIP. Lorsqu'une question est posée, EduAI recherche les trois passages les plus pertinents, les injecte dans le prompt, puis Llama 3.3 70B génère une réponse contextualisée. Le même socle permet ensuite de produire résumés, quiz, flashcards, mindmap et examens.
""",
    )


def slide_9(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Architecture", "Full-stack moderne, modulaire et déployable", "Une architecture compatible produit SaaS : web, API, données, IA, conteneurs.", 9, CARD)
    layers = [
        ("Frontend", "Next.js 14 - React 18 - TypeScript - Tailwind - Framer Motion - ReactFlow - Recharts", BLUE_SOFT, BLUE),
        ("Backend API", "FastAPI - Python - Pydantic - Motor - routers par domaine", COPPER_SOFT, COPPER),
        ("IA/NLP", "PyMuPDF - MiniLM - FAISS - spaCy - CamemBERT fallback - Groq Llama 3.3 70B", GREEN_SOFT, GREEN),
        ("Données", "MongoDB - collections users/courses/chunks/summaries/quizzes/flashcards/exams", SUBTLE, NAVY),
        ("Déploiement", "Docker Compose - Portainer - volumes persistants - health checks", C("#F0E3DA"), COPPER_DARK),
    ]
    for i, (title, body, bg, col) in enumerate(layers):
        y = 1.95 + i * 0.88
        rect(s, 0.95, y, 11.45, 0.66, bg, BORDER)
        text(s, title, 1.22, y + 0.18, 1.8, 0.16, 11.5, col, True)
        text(s, body, 3.05, y + 0.18, 8.9, 0.16, 10.3, INK_2)
    add_notes(
        s,
        """
L'architecture a été pensée pour être modulaire. Le frontend est développé avec Next.js, React et TypeScript. Le backend repose sur FastAPI, ce qui est adapté à Python et aux traitements IA. MongoDB stocke les utilisateurs, les cours et les résultats générés, tandis que FAISS gère la recherche vectorielle. Le déploiement est prévu avec Docker Compose et Portainer, avec des volumes persistants et des vérifications de santé. L'objectif est d'avoir une base fonctionnelle, maintenable et déployable.
""",
    )


def slide_10(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Conception", "Preuves issues du workspace : UML, séquences et MLD", "Le projet a été modélisé, implémenté et documenté, pas seulement codé rapidement.", 10)
    picture_fit(s, FIG / "uc_global.png", 0.85, 2.0, 2.5, 3.4)
    picture_fit(s, FIG / "sd_s2_upload.png", 3.65, 2.0, 4.15, 2.0)
    picture_fit(s, FIG / "sd_s3_qa.png", 8.1, 2.0, 4.15, 2.0)
    picture_fit(s, FIG / "mld_2_quiz_flash.png", 3.65, 4.35, 4.15, 1.35)
    picture_fit(s, FIG / "sd_s9_admin.png", 8.1, 4.35, 4.15, 1.35)
    chip(s, "Cas d'utilisation", 1.1, 5.65, 1.7, NAVY, size=7.5)
    chip(s, "Upload PDF", 4.8, 4.05, 1.25, COPPER, size=7.5)
    chip(s, "Q&A RAG", 9.6, 4.05, 1.05, BLUE, size=7.5)
    chip(s, "Données", 5.1, 5.75, 1.0, GREEN, size=7.5)
    chip(s, "Admin", 9.75, 5.75, 0.9, NAVY, size=7.5)
    add_notes(
        s,
        """
Pour donner de la crédibilité au projet, je ne me suis pas limité à l'implémentation. J'ai aussi modélisé les cas d'utilisation, les séquences principales et le modèle logique de données. Ici, on voit par exemple le cas global, la séquence d'upload et d'indexation, la séquence Q&A RAG, une partie du modèle de données et la partie administration. Cela montre que l'application a été pensée comme un système complet.
""",
    )


def slide_11(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Crédibilité", "Validation, sécurité et robustesse", "La soutenance doit prouver que le projet fonctionne et que les risques sont maîtrisés.", 11, CARD)
    modules = ["Auth JWT", "Upload + FAISS", "Q&A SSE", "Résumé cache", "Quiz scoring", "Flashcards SM-2", "Mindmap", "Examen A-F", "Admin rôle"]
    for i, m in enumerate(modules):
        x = 0.85 + (i % 3) * 3.9
        y = 2.0 + (i // 3) * 0.78
        rect(s, x, y, 3.45, 0.55, GREEN_SOFT, C("#CFE8D9"))
        text(s, "✓", x + 0.22, y + 0.16, 0.2, 0.1, 11, GREEN, True, align=PP_ALIGN.CENTER)
        text(s, m, x + 0.55, y + 0.15, 2.45, 0.12, 9.2, INK, True)
    blocks = [
        ("Sécurité", "JWT HS256, bcrypt, rôles user/admin, validation Pydantic.", NAVY),
        ("Robustesse IA", "Sources RAG, fallback CamemBERT, cache MongoDB.", BLUE),
        ("Déploiement", "Docker Compose, Portainer, volumes persistants, health checks.", COPPER),
    ]
    for i, (t, b, c) in enumerate(blocks):
        mini_card(s, t, b, 0.85 + i * 3.9, 5.05, 3.45, 1.0, c)
    add_notes(
        s,
        """
Pour la validation, j'ai testé les modules principaux à plusieurs niveaux : tests unitaires pour les services, tests d'intégration pour les endpoints, scripts de bout en bout pour la chaîne IA, et tests manuels sur des scénarios utilisateur réels. Les modules prévus dans le cahier des charges sont couverts : authentification, upload, RAG, résumé, quiz, flashcards, mindmap, examen et administration. La sécurité repose sur JWT, bcrypt, les rôles et la validation des données.
""",
    )


def slide_12(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Positionnement", "EduAI combine ce que les autres outils séparent", "La différenciation vient de l'intégration autour des vrais cours PDF.", 12)
    headers = ["Solution", "PDF perso", "RAG", "Quiz", "Flashcards", "Examen", "Prix adapté"]
    rows = [
        ["ChatGPT", "△", "△", "△", "○", "○", "△"],
        ["Quizlet / Anki", "○", "○", "✓", "✓", "○", "△"],
        ["Notion AI / SciSpace", "△", "△", "○", "○", "○", "△"],
        ["Coursera / EdX", "○", "○", "✓", "○", "✓", "○"],
        ["EduAI", "✓", "✓", "✓", "✓", "✓", "✓"],
    ]
    x0, y0 = 0.7, 2.05
    widths = [2.15, 1.35, 0.9, 0.9, 1.35, 1.05, 1.45]
    for c, h in enumerate(headers):
        x = x0 + sum(widths[:c])
        rect(s, x, y0, widths[c], 0.5, NAVY, None, radius=False)
        text(s, h, x + 0.05, y0 + 0.17, widths[c] - 0.1, 0.1, 7.4, WHITE, True, align=PP_ALIGN.CENTER)
    for r, row in enumerate(rows):
        y = y0 + 0.5 * (r + 1)
        is_edu = row[0] == "EduAI"
        for c, val in enumerate(row):
            x = x0 + sum(widths[:c])
            rect(s, x, y, widths[c], 0.5, BLUE_SOFT if is_edu else CARD, BORDER, radius=False)
            text(s, val, x + 0.05, y + 0.17, widths[c] - 0.1, 0.1, 11 if c else 8.5, BLUE if is_edu else INK_2, is_edu or c == 0, align=PP_ALIGN.CENTER)
    rect(s, 0.9, 5.45, 11.4, 0.68, COPPER_SOFT, C("#ECD4C5"))
    text(s, "Positionnement : EduAI n'est pas un chatbot générique ; c'est une plateforme de révision contextualisée.", 1.25, 5.68, 10.7, 0.12, 11, COPPER_DARK, True, align=PP_ALIGN.CENTER)
    add_notes(
        s,
        """
Les solutions existantes répondent chacune à une partie du problème. ChatGPT est puissant, mais il reste généraliste. Anki et Quizlet sont utiles, mais demandent souvent une création manuelle. Notion AI ou SciSpace peuvent aider à résumer, mais ils ne couvrent pas tout le cycle d'apprentissage. La différence d'EduAI est l'intégration : un cours PDF réel, traité automatiquement, puis transformé en plusieurs outils pédagogiques dans une seule plateforme, avec une logique adaptée aux étudiants marocains.
""",
    )


def slide_13(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Business", "Un SaaS freemium accessible, avec potentiel B2B", "Des hypothèses progressives, un coût de lancement maîtrisé et une trajectoire réaliste.", 13, CARD)
    metric(s, "0 MAD", "Free", "Jusqu'à 3 cours, découverte de la valeur produit.", 0.85, 2.05, 2.75, BLUE)
    metric(s, "49 MAD", "Pro / mois", "Accès illimité, examen, mindmap, export PDF, quiz avancés.", 3.9, 2.05, 2.75, COPPER)
    rect(s, 7.1, 2.05, 5.15, 2.2, SUBTLE, BORDER)
    text(s, "Revenus annuels projetés", 7.42, 2.35, 2.8, 0.18, 13, INK, True)
    vals = [147000, 940800, 4233600]
    labels = ["A1", "A2", "A3"]
    names = ["147k", "940,8k", "4,23M"]
    for i, (v, lab, name, col) in enumerate(zip(vals, labels, names, [BLUE, COPPER, GREEN])):
        h = 1.18 * v / max(vals)
        x = 7.55 + i * 1.25
        rect(s, x, 3.75 - h, 0.58, h, col, None, radius=False)
        text(s, name, x - 0.15, 3.52 - h, 0.9, 0.1, 7.3, col, True, align=PP_ALIGN.CENTER)
        text(s, lab, x, 3.9, 0.58, 0.1, 7.4, INK_2, True, align=PP_ALIGN.CENTER)
    text(s, "MAD/an", 11.25, 3.72, 0.55, 0.1, 7.3, MUTED)
    rect(s, 0.85, 4.75, 5.8, 0.85, NAVY, None)
    text(s, "Coût mensuel estimé au lancement", 1.18, 5.05, 3.55, 0.12, 10.5, C("#FBD6BF"), True)
    text(s, "1 300 MAD", 4.72, 4.96, 1.35, 0.22, 18, WHITE, True, align=PP_ALIGN.RIGHT)
    rect(s, 7.1, 4.75, 5.15, 0.85, GREEN_SOFT, C("#CFE8D9"))
    text(s, "Perspective B2B", 7.42, 5.05, 1.7, 0.12, 10.5, GREEN, True)
    text(s, "écoles, universités, centres de formation", 9.0, 5.05, 2.75, 0.12, 10, INK_2)
    text(s, "Important : présenter ces chiffres comme des projections, pas comme des garanties.", 0.9, 6.36, 7.3, 0.1, 8.2, MUTED)
    add_notes(
        s,
        """
Le modèle économique choisi est un freemium adapté au pouvoir d'achat des étudiants marocains. Le plan gratuit permet de découvrir la plateforme, tandis que le plan Pro à 49 MAD par mois donne accès aux fonctions avancées. Les projections restent volontairement progressives : 5 000 utilisateurs en première année, puis une croissance organique. Le coût de lancement estimé est autour de 1 300 MAD par mois, ce qui rend l'expérimentation réaliste. À moyen terme, une piste importante est le B2B avec des écoles ou universités.
""",
    )


def slide_14(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Impact", "Valeur pédagogique, sociale et économique", "EduAI aide sans remplacer : il accompagne l'étudiant entre le cours et l'examen.", 14)
    impacts = [
        ("Pédagogique", "Comprendre les chapitres, réviser activement, mesurer ses acquis.", BLUE),
        ("Social", "Rendre l'aide intelligente plus accessible aux étudiants marocains.", COPPER),
        ("Technique", "Construire une base IA/NLP déployable et maintenable.", GREEN),
        ("Économique", "Évoluer vers SaaS freemium puis offres institutionnelles.", NAVY),
    ]
    for i, (t, b, c) in enumerate(impacts):
        x = 0.9 + (i % 2) * 5.8
        y = 2.05 + (i // 2) * 1.55
        mini_card(s, t, b, x, y, 5.25, 1.15, c)
    rect(s, 0.9, 5.55, 11.4, 0.66, CARD, BORDER)
    text(s, "Perspectives : application mobile, paiements locaux CMI/Payzone, OCR pour PDF scannés, pilote B2B avec établissements.", 1.22, 5.78, 10.75, 0.12, 10, INK_2, True, align=PP_ALIGN.CENTER)
    add_notes(
        s,
        """
L'impact attendu d'EduAI se situe à plusieurs niveaux. Sur le plan pédagogique, l'étudiant comprend mieux, révise activement et s'auto-évalue avant l'examen. Sur le plan social, la solution vise une accessibilité adaptée au contexte marocain. Sur le plan technique, le projet démontre la capacité à construire un système IA complet. Sur le plan économique, il peut évoluer vers un SaaS freemium puis vers des offres institutionnelles. Les perspectives incluent le mobile, les paiements locaux, l'OCR et les partenariats B2B.
""",
    )


def slide_15(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, NAVY)
    rect(s, 0, 0, 13.333, 7.5, NAVY, None, radius=False)
    brand(s, 0.65, 0.55, dark=True)
    text(s, "CONCLUSION", 0.72, 1.38, 2.4, 0.18, 9, C("#FBD6BF"), True, "JetBrains Mono")
    text(s, "EduAI ne remplace pas l'enseignant.", 1.05, 2.18, 11.2, 0.42, 29, WHITE, True, "Playfair Display", PP_ALIGN.CENTER)
    text(
        s,
        "Il donne à chaque étudiant un compagnon de révision intelligent, accessible et contextualisé.",
        1.85,
        3.08,
        9.65,
        0.78,
        23,
        C("#FDBA74"),
        True,
        "DM Sans",
        PP_ALIGN.CENTER,
    )
    chip(s, "Comprendre mieux", 2.0, 4.85, 2.35, BLUE, size=11)
    chip(s, "Réviser efficacement", 5.15, 4.85, 2.75, COPPER, size=11)
    chip(s, "S'auto-évaluer", 8.7, 4.85, 2.35, GREEN, size=11)
    text(s, "Merci pour votre attention", 4.35, 6.25, 4.65, 0.22, 15, C("#D7CCC4"), True, align=PP_ALIGN.CENTER)
    footer(s, 15, dark=True)
    add_notes(
        s,
        """
Pour conclure, EduAI n'a pas vocation à remplacer l'enseignant. Le rôle de l'enseignant reste central. EduAI intervient plutôt entre le cours et l'examen, au moment où l'étudiant doit comprendre, réviser et vérifier son niveau. Ce projet m'a permis de construire une solution complète qui combine développement web, NLP, RAG, LLM, déploiement et réflexion business. Aujourd'hui, EduAI est une base fonctionnelle, testée et déployable, avec un potentiel d'évolution vers un produit EdTech adapté au contexte marocain. Je vous remercie pour votre attention.
""",
    )


def build():
    prs = Presentation()
    prs.slide_width = I(SLIDE_W)
    prs.slide_height = I(SLIDE_H)

    for fn in [
        slide_1,
        slide_2,
        slide_3,
        slide_4,
        slide_5,
        slide_6,
        slide_7,
        slide_8,
        slide_9,
        slide_10,
        slide_11,
        slide_12,
        slide_13,
        slide_14,
        slide_15,
    ]:
        fn(prs)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Generated: {OUT}")
    print(f"Slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()

